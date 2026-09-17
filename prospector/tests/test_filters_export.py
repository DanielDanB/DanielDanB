import io

from openpyxl import load_workbook

from app.export import do_csv, do_xlsx
from app.filters import Filtr, hledej, spocitej
from app.storage import je_na_blacklistu, pridej_na_blacklist, uloz_kontakty


def nazvy(conn, filtr):
    return [r["nazev"] for r in hledej(conn, filtr)]


def test_vychozi_filtr_vylouci_nespolehliveho_platce(naplnena):
    assert nazvy(naplnena, Filtr()) == ["Alfa CNC s.r.o."]
    assert "Beta Velkoobchod a.s." in nazvy(naplnena, Filtr(vyloucit_nespolehlive=False))


def test_nace_prefix_chyti_i_vedlejsi_kod(naplnena):
    # 4690 je u Bety az druhy v poradi
    assert nazvy(naplnena, Filtr(nace=["4690"], vyloucit_nespolehlive=False)) \
        == ["Beta Velkoobchod a.s."]


def test_rozsah_zamestnancu_bere_prekryv_pasma(naplnena):
    assert nazvy(naplnena, Filtr(zamestnanci_min=30)) == ["Alfa CNC s.r.o."]
    assert nazvy(naplnena, Filtr(zamestnanci_min=200)) == []
    # pasmo 25-49 se s rozsahem 10-30 prekryva, takze vyhovuje
    assert nazvy(naplnena, Filtr(zamestnanci_min=10, zamestnanci_max=30)) == ["Alfa CNC s.r.o."]


def test_filtr_kraje_a_webu(naplnena):
    assert nazvy(naplnena, Filtr(kraje=["CZ064"])) == ["Alfa CNC s.r.o."]
    assert nazvy(naplnena, Filtr(kraje=["CZ041"])) == []
    assert nazvy(naplnena, Filtr(jen_s_webem=True)) == ["Alfa CNC s.r.o."]


def test_firma_na_blacklistu_zmizi_z_vysledku(naplnena):
    pridej_na_blacklist(naplnena, "00000001", "ico", "nepřeje si oslovení")
    assert nazvy(naplnena, Filtr()) == []
    assert spocitej(naplnena, Filtr()) == 0


def test_blacklist_domeny_blokuje_i_adresy_pod_ni(naplnena):
    pridej_na_blacklist(naplnena, "alfa.cz", "domena", "opt-out")
    assert je_na_blacklistu(naplnena, "cokoliv@alfa.cz")
    pridano = uloz_kontakty(naplnena, [{"ico": "00000001", "typ": "email",
                                        "hodnota": "info@alfa.cz", "role": "obecny"}])
    assert pridano == 0


def test_csv_obsahuje_kontakt_ve_spravnem_sloupci(naplnena):
    obsah = do_csv(naplnena, Filtr())
    assert obsah.startswith("﻿")  # BOM kvuli Excelu
    radky = obsah.replace("﻿", "").strip().splitlines()
    hlavicka, prvni = radky[0].split(";"), radky[1].split(";")
    assert prvni[hlavicka.index("E-mail nákup")] == "nakup@alfa.cz"
    assert prvni[hlavicka.index("Statutární orgán")] == "Jan Novák (jednatel)"


def test_xlsx_ma_list_se_zdroji(naplnena):
    sesit = load_workbook(io.BytesIO(do_xlsx(naplnena, Filtr())))
    assert sesit.sheetnames == ["Firmy", "Kontakty a zdroje"]
    zdroje = list(sesit["Kontakty a zdroje"].values)
    assert zdroje[1][3] == "nakup@alfa.cz"
    assert zdroje[1][6] == "https://alfa.cz/kontakty"  # provenience je v exportu
